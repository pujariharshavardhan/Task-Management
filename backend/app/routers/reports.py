"""Reports (8 types from the spec), PDF/Excel/CSV export, and rule-based risk detection."""
import csv
import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..deps import accessible_project_ids, get_current_user
from ..models import Milestone, Project, Task, User
from ..services import INACTIVE_STATUSES, is_at_risk

router = APIRouter(tags=["Reports"])

REPORT_TYPES = {
    "project-progress": "Project Progress",
    "task-completion": "Task Completion",
    "team-workload": "Team Workload",
    "delayed-tasks": "Delayed Tasks",
    "deadline-report": "Deadline Report",
    "milestone-report": "Milestone Report",
    "member-contribution": "Member Contribution",
    "project-completion": "Project Completion Report",
}


def _scope(db: Session, user: User, project_id: int | None):
    pids = accessible_project_ids(db, user)
    if project_id is not None:
        if project_id not in pids:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Project not found")
        pids = [project_id]
    projects = db.query(Project).filter(Project.id.in_(pids), Project.is_archived.is_(False)).order_by(Project.end_date).all()
    ids = [p.id for p in projects]
    tasks = db.query(Task).filter(Task.project_id.in_(ids)).options(selectinload(Task.project), selectinload(Task.assignee)).all()
    milestones = db.query(Milestone).filter(Milestone.project_id.in_(ids)).options(selectinload(Milestone.project)).order_by(Milestone.due_date).all()
    return projects, tasks, milestones


def _metrics(projects, tasks, today):
    done = [t for t in tasks if t.status == "Done"]
    delayed = [t for t in tasks if t.status != "Done" and t.due_date and t.due_date < today]
    on_time = [t for t in done if not t.due_date or (t.updated_at and t.updated_at.date() <= t.due_date)]
    return {
        "overallProgress": round(sum(p.progress for p in projects) / len(projects)) if projects else 0,
        "totalTasks": len(tasks), "completedTasks": len(done), "pendingTasks": len(tasks) - len(done),
        "onTimeTasks": len(on_time), "delayedTasks": len(delayed),
    }


def _workload(tasks):
    people: dict[int, dict] = {}
    for t in tasks:
        if not t.assignee:
            continue
        w = people.setdefault(t.assigned_to, {"member": t.assignee.name, "userId": t.assigned_to, "total": 0, "Todo": 0, "In Progress": 0, "Review": 0, "Done": 0, "estimatedHours": 0.0, "actualHours": 0.0})
        w["total"] += 1
        w[t.status] += 1
        w["estimatedHours"] += t.estimated_hours or 0
        w["actualHours"] += t.actual_hours or 0
    return sorted(people.values(), key=lambda w: -w["total"])


def build_report(kind: str, projects, tasks, milestones) -> dict:
    """Returns {title, columns, rows, metrics, chart} — rows are plain values so they export cleanly."""
    today = date.today()
    title = REPORT_TYPES[kind]
    chart = None
    if kind in ("project-progress", "project-completion"):
        cols = ["Project", "Code", "Status", "Progress %", "Total Tasks", "Completed", "Start", "Deadline", "Completed On", "At Risk"]
        rows = []
        for p in projects:
            pt = [t for t in tasks if t.project_id == p.id]
            rows.append([p.name, p.project_code, p.status, p.progress, len(pt), sum(t.status == "Done" for t in pt), p.start_date, p.end_date, p.actual_completion_date, "Yes" if is_at_risk(p, today) else "No"])
        if kind == "project-completion":
            rows = [r for r, p in zip(rows, projects) if p.status == "Completed" or p.progress >= 100] or rows
        chart = [{"name": p.project_code, "label": p.name, "progress": p.progress} for p in projects]
    elif kind == "task-completion":
        cols = ["Task ID", "Task", "Project", "Assignee", "Status", "Priority", "Due", "Est. Hours", "Actual Hours"]
        rows = [[t.task_key, t.title, t.project.name, t.assignee.name if t.assignee else "", t.status, t.priority, t.due_date, t.estimated_hours, t.actual_hours] for t in tasks]
        chart = [{"name": s, "value": sum(t.status == s for t in tasks)} for s in ["Todo", "In Progress", "Review", "Done"]]
    elif kind in ("team-workload", "member-contribution"):
        cols = ["Member", "Total", "Todo", "In Progress", "Review", "Done", "Est. Hours", "Actual Hours", "Completion %"]
        wl = _workload(tasks)
        rows = [[w["member"], w["total"], w["Todo"], w["In Progress"], w["Review"], w["Done"], round(w["estimatedHours"], 1), round(w["actualHours"], 1), round(w["Done"] * 100 / w["total"]) if w["total"] else 0] for w in wl]
        chart = wl
    elif kind == "delayed-tasks":
        cols = ["Task ID", "Task", "Project", "Assignee", "Due", "Days Overdue", "Priority"]
        late = sorted([t for t in tasks if t.status != "Done" and t.due_date and t.due_date < today], key=lambda t: t.due_date)
        rows = [[t.task_key, t.title, t.project.name, t.assignee.name if t.assignee else "", t.due_date, (today - t.due_date).days, t.priority] for t in late]
    elif kind == "deadline-report":
        cols = ["Task ID", "Task", "Project", "Assignee", "Due", "Days Left", "Priority"]
        soon = sorted([t for t in tasks if t.status != "Done" and t.due_date and t.due_date >= today], key=lambda t: t.due_date)
        rows = [[t.task_key, t.title, t.project.name, t.assignee.name if t.assignee else "", t.due_date, (t.due_date - today).days, t.priority] for t in soon]
    elif kind == "milestone-report":
        cols = ["Milestone", "Project", "Start", "Due", "Progress %", "Status"]
        rows = [[m.name, m.project.name, m.start_date, m.due_date, m.progress, m.status] for m in milestones]
        chart = [{"name": m.name, "progress": m.progress} for m in milestones]
    else:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown report type")
    return {"type": kind, "title": title, "generatedOn": today, "metrics": _metrics(projects, tasks, today), "columns": cols, "rows": rows, "chart": chart}


@router.get("/reports", summary="Available report types")
def report_types():
    return [{"key": k, "title": v} for k, v in REPORT_TYPES.items()]


@router.get("/reports/{kind}")
def get_report(kind: str, project_id: int | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if kind not in REPORT_TYPES:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown report type")
    return build_report(kind, *_scope(db, user, project_id))


def _fmt(v):
    if v is None:
        return ""
    return v.isoformat() if isinstance(v, date) else v


@router.get("/reports/{kind}/export", summary="Download a report as csv, xlsx or pdf")
def export_report(kind: str, format: str = Query("csv", pattern="^(csv|xlsx|pdf)$"), project_id: int | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if kind not in REPORT_TYPES:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown report type")
    rep = build_report(kind, *_scope(db, user, project_id))
    fname = f"{kind}-{date.today().isoformat()}.{format}"
    rows = [[_fmt(v) for v in r] for r in rep["rows"]]

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(rep["columns"])
        w.writerows(rows)
        data, media = buf.getvalue().encode("utf-8-sig"), "text/csv"
    elif format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = rep["title"][:31]
        ws.append(rep["columns"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0A0A0A")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(45, max(len(str(c.value or "")) for c in col) + 3)
        out = io.BytesIO()
        wb.save(out)
        data, media = out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28, title=rep["title"])
        styles = getSampleStyleSheet()
        m = rep["metrics"]
        story = [
            Paragraph(f"APMS — {rep['title']}", styles["Title"]),
            Paragraph(f"Generated {date.today():%B %d, %Y} for {user.name}", styles["Normal"]),
            Spacer(1, 8),
            Paragraph(f"Overall progress {m['overallProgress']}% · Tasks {m['totalTasks']} · Completed {m['completedTasks']} · Pending {m['pendingTasks']} · On-time {m['onTimeTasks']} · Delayed {m['delayedTasks']}", styles["Normal"]),
            Spacer(1, 12),
        ]
        cell = styles["BodyText"]
        cell.fontSize = 8
        table = Table([rep["columns"]] + [[Paragraph(str(v), cell) for v in r] for r in rows] if rows else [rep["columns"], ["No data"] + [""] * (len(rep["columns"]) - 1)], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0A0A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, 0), 9), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F9")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(table)
        doc.build(story)
        data, media = out.getvalue(), "application/pdf"
    return StreamingResponse(io.BytesIO(data), media_type=media, headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/insights/risks", summary="Risk & deadline detection: overdue work, overloaded members, at-risk projects, blocked reviews")
def risks(project_id: int | None = None, overload_threshold: int = 4, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    projects, tasks, milestones = _scope(db, user, project_id)
    today = date.today()
    open_tasks = [t for t in tasks if t.status != "Done"]
    findings = []
    for p in projects:
        if is_at_risk(p, today):
            findings.append({"severity": "high", "kind": "project_at_risk", "projectId": p.id, "message": f"{p.name} is {p.progress}% done with {(p.end_date - today).days} days left"})
        elif p.status not in INACTIVE_STATUSES and p.end_date and 0 <= (p.end_date - today).days <= 7:
            findings.append({"severity": "medium", "kind": "project_deadline", "projectId": p.id, "message": f"{p.name} is due in {(p.end_date - today).days} days"})
    for t in open_tasks:
        if t.due_date and t.due_date < today:
            findings.append({"severity": "high" if t.priority in ("High", "Critical") else "medium", "kind": "overdue_task", "projectId": t.project_id, "taskId": t.id, "message": f'"{t.title}" is {(today - t.due_date).days} days overdue'})
        elif t.due_date and (t.due_date - today).days <= 2:
            findings.append({"severity": "low", "kind": "deadline_soon", "projectId": t.project_id, "taskId": t.id, "message": f'"{t.title}" is due {"today" if t.due_date == today else f"in {(t.due_date - today).days} days"}'})
        if t.status == "Review" and t.updated_at and (today - t.updated_at.date()).days >= 3:
            findings.append({"severity": "medium", "kind": "blocked_review", "projectId": t.project_id, "taskId": t.id, "message": f'"{t.title}" has waited in Review for {(today - t.updated_at.date()).days} days'})
        if t.estimated_hours and t.actual_hours and t.actual_hours > t.estimated_hours * 1.25:
            findings.append({"severity": "low", "kind": "over_estimate", "projectId": t.project_id, "taskId": t.id, "message": f'"{t.title}" is over its estimate ({t.actual_hours:g}h of {t.estimated_hours:g}h)'})
    for w in _workload(open_tasks):
        if w["total"] >= overload_threshold:
            findings.append({"severity": "medium", "kind": "overloaded_member", "userId": w["userId"], "message": f"{w['member']} has {w['total']} open tasks"})
    for m in milestones:
        if m.progress < 100 and m.due_date and m.due_date < today:
            findings.append({"severity": "high", "kind": "milestone_overdue", "projectId": m.project_id, "message": f"Milestone {m.name} ({m.project.name}) is overdue at {m.progress}%"})
    order = {"high": 0, "medium": 1, "low": 2}
    findings.sort(key=lambda f: order[f["severity"]])
    return {"generatedOn": today, "counts": {s: sum(f["severity"] == s for f in findings) for s in order}, "findings": findings}
